import streamlit as st
import time
import io
from my_agents import build_reader_agent, build_search_agent, build_writer_chain, critic_chain

# ── Page config ───────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="ResearchOrbit · AI Research Agent",
    page_icon="🔬",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ── CSS ────────────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Instrument+Serif:ital@0;1&family=DM+Mono:wght@300;400;500&family=Outfit:wght@300;400;500;600&display=swap');

html, body, [class*="css"] {
    font-family: 'Outfit', sans-serif;
    color: #e8e3d8;
}

.stApp {
    background: #080b10;
    background-image:
        radial-gradient(ellipse 70% 55% at 15% -5%,  rgba(99,179,237,0.07) 0%, transparent 60%),
        radial-gradient(ellipse 55% 45% at 85% 105%, rgba(159,122,234,0.06) 0%, transparent 55%),
        radial-gradient(ellipse 40% 30% at 50%  50%, rgba(49,151,149,0.04) 0%, transparent 60%);
}

#MainMenu, footer, header { visibility: hidden; }
.block-container { padding: 2rem 3rem 5rem; max-width: 1300px; }

/* ── Hero ── */
.hero { text-align: center; padding: 3rem 0 2rem; }
.hero-badge {
    display: inline-block;
    font-family: 'DM Mono', monospace;
    font-size: 0.65rem;
    letter-spacing: 0.22em;
    text-transform: uppercase;
    color: #63b3ed;
    border: 1px solid rgba(99,179,237,0.25);
    border-radius: 20px;
    padding: 0.3rem 0.9rem;
    margin-bottom: 1.4rem;
    opacity: 0.85;
}
.hero h1 {
    font-family: 'Instrument Serif', serif;
    font-size: clamp(3rem, 7vw, 5.5rem);
    font-weight: 400;
    font-style: italic;
    line-height: 1.0;
    letter-spacing: -0.02em;
    color: #f0ebe0;
    margin: 0 0 1rem;
}
.hero h1 em { color: #63b3ed; font-style: normal; }
.hero-sub {
    font-size: 1rem;
    font-weight: 300;
    color: #7a8694;
    max-width: 500px;
    margin: 0 auto;
    line-height: 1.7;
}
.rule {
    height: 1px;
    background: linear-gradient(90deg, transparent, rgba(99,179,237,0.2), rgba(159,122,234,0.2), transparent);
    margin: 2rem 0;
}

/* ── Cards ── */
.card {
    background: rgba(255,255,255,0.025);
    border: 1px solid rgba(255,255,255,0.07);
    border-radius: 16px;
    padding: 1.8rem 2rem;
    margin-bottom: 1.2rem;
}
.card-blue  { border-color: rgba(99,179,237,0.18);  background: rgba(99,179,237,0.04); }
.card-green { border-color: rgba(72,187,120,0.18);  background: rgba(72,187,120,0.04); }

/* ── Inputs ── */
.stTextInput > div > div > input {
    background: rgba(255,255,255,0.04) !important;
    border: 1px solid rgba(255,255,255,0.1) !important;
    border-radius: 10px !important;
    color: #f0ebe0 !important;
    font-family: 'Outfit', sans-serif !important;
    font-size: 1rem !important;
    padding: 0.75rem 1rem !important;
    transition: border-color 0.2s, box-shadow 0.2s !important;
}
.stTextInput > div > div > input:focus {
    border-color: rgba(99,179,237,0.5) !important;
    box-shadow: 0 0 0 3px rgba(99,179,237,0.1) !important;
}
.stTextInput > label {
    font-family: 'DM Mono', monospace !important;
    font-size: 0.68rem !important;
    letter-spacing: 0.18em !important;
    text-transform: uppercase !important;
    color: #63b3ed !important;
    font-weight: 500 !important;
}

/* ── Selects ── */
.stSelectbox > label {
    font-family: 'DM Mono', monospace !important;
    font-size: 0.68rem !important;
    letter-spacing: 0.18em !important;
    text-transform: uppercase !important;
    color: #9f7aea !important;
    font-weight: 500 !important;
}

/* ── Button ── */
.stButton > button {
    background: linear-gradient(135deg, #2b6cb0 0%, #553c9a 100%) !important;
    color: #fff !important;
    font-family: 'Outfit', sans-serif !important;
    font-weight: 600 !important;
    font-size: 0.92rem !important;
    border: none !important;
    border-radius: 10px !important;
    padding: 0.7rem 2rem !important;
    transition: transform 0.15s, box-shadow 0.15s !important;
    box-shadow: 0 4px 20px rgba(43,108,176,0.3) !important;
    width: 100%;
}
.stButton > button:hover {
    transform: translateY(-2px) !important;
    box-shadow: 0 8px 28px rgba(43,108,176,0.45) !important;
}

/* ── Pipeline steps ── */
.step {
    display: flex;
    align-items: center;
    gap: 1rem;
    padding: 1rem 1.2rem;
    border-radius: 12px;
    border: 1px solid rgba(255,255,255,0.05);
    margin-bottom: 0.7rem;
    background: rgba(255,255,255,0.02);
    transition: all 0.3s ease;
    position: relative;
    overflow: hidden;
}
.step.active {
    border-color: rgba(99,179,237,0.35);
    background: rgba(99,179,237,0.05);
}
.step.active::before {
    content: '';
    position: absolute;
    left: 0; top: 0; bottom: 0;
    width: 3px;
    background: linear-gradient(180deg, #63b3ed, #9f7aea);
    border-radius: 12px 0 0 12px;
}
.step.done {
    border-color: rgba(72,187,120,0.3);
    background: rgba(72,187,120,0.04);
}
.step.done::before {
    content: '';
    position: absolute;
    left: 0; top: 0; bottom: 0;
    width: 3px;
    background: #48bb78;
    border-radius: 12px 0 0 12px;
}
.step-icon {
    font-size: 1.2rem;
    width: 36px;
    height: 36px;
    display: flex;
    align-items: center;
    justify-content: center;
    border-radius: 8px;
    background: rgba(255,255,255,0.05);
    flex-shrink: 0;
}
.step-body { flex: 1; }
.step-title {
    font-family: 'Outfit', sans-serif;
    font-size: 0.88rem;
    font-weight: 600;
    color: #e8e3d8;
    margin-bottom: 0.1rem;
}
.step-desc {
    font-size: 0.75rem;
    color: #536070;
}
.step-badge {
    font-family: 'DM Mono', monospace;
    font-size: 0.62rem;
    letter-spacing: 0.1em;
    padding: 0.2rem 0.55rem;
    border-radius: 6px;
    flex-shrink: 0;
}
.badge-wait   { color: #3d4a56; background: rgba(255,255,255,0.03); border: 1px solid rgba(255,255,255,0.05); }
.badge-run    { color: #63b3ed; background: rgba(99,179,237,0.1);   border: 1px solid rgba(99,179,237,0.2); }
.badge-done   { color: #48bb78; background: rgba(72,187,120,0.1);   border: 1px solid rgba(72,187,120,0.2); }

/* ── Results ── */
.result-box {
    background: rgba(255,255,255,0.02);
    border: 1px solid rgba(255,255,255,0.07);
    border-radius: 14px;
    padding: 1.5rem 1.8rem;
    margin-bottom: 1rem;
}
.result-box pre {
    font-family: 'DM Mono', monospace;
    font-size: 0.8rem;
    color: #7a8694;
    white-space: pre-wrap;
    margin: 0;
}
.section-label {
    font-family: 'DM Mono', monospace;
    font-size: 0.65rem;
    letter-spacing: 0.22em;
    text-transform: uppercase;
    margin-bottom: 0.9rem;
    padding-bottom: 0.6rem;
    border-bottom: 1px solid rgba(255,255,255,0.06);
}
.label-blue   { color: #63b3ed; }
.label-green  { color: #48bb78; }
.label-purple { color: #9f7aea; }

/* ── Export row ── */
.export-row {
    display: flex;
    gap: 0.8rem;
    flex-wrap: wrap;
    margin-top: 1rem;
}

/* ── Download buttons ── */
.stDownloadButton > button {
    background: rgba(255,255,255,0.04) !important;
    color: #c8d3de !important;
    font-family: 'DM Mono', monospace !important;
    font-size: 0.75rem !important;
    letter-spacing: 0.05em !important;
    border: 1px solid rgba(255,255,255,0.1) !important;
    border-radius: 8px !important;
    padding: 0.5rem 1.2rem !important;
    transition: all 0.15s !important;
    box-shadow: none !important;
    width: auto !important;
}
.stDownloadButton > button:hover {
    background: rgba(255,255,255,0.07) !important;
    border-color: rgba(99,179,237,0.35) !important;
    color: #63b3ed !important;
    transform: translateY(-1px) !important;
}

/* ── Score badge ── */
.score-chip {
    display: inline-block;
    font-family: 'Instrument Serif', serif;
    font-style: italic;
    font-size: 2.2rem;
    color: #63b3ed;
    margin-bottom: 0.5rem;
}

/* ── Settings labels ── */
.settings-heading {
    font-family: 'DM Mono', monospace;
    font-size: 0.65rem;
    letter-spacing: 0.2em;
    text-transform: uppercase;
    color: #536070;
    margin-bottom: 0.6rem;
}

/* ── Footer ── */
.footer {
    font-family: 'DM Mono', monospace;
    font-size: 0.68rem;
    color: #2e3a45;
    text-align: center;
    margin-top: 4rem;
    letter-spacing: 0.08em;
}

/* ── Spinner ── */
.stSpinner > div { color: #63b3ed !important; }

/* ── Expander ── */
details summary {
    font-family: 'DM Mono', monospace !important;
    font-size: 0.72rem !important;
    color: #536070 !important;
    letter-spacing: 0.1em !important;
}
</style>
""", unsafe_allow_html=True)


# ── Helpers ───────────────────────────────────────────────────────────────────
def render_step(icon: str, title: str, desc: str, state: str):
    state_map = {
        "waiting": ("WAITING", "badge-wait", ""),
        "running": ("● RUNNING", "badge-run", "active"),
        "done":    ("✓ DONE",   "badge-done", "done"),
    }
    label, badge_cls, card_cls = state_map.get(state, ("", "", ""))
    st.markdown(f"""
    <div class="step {card_cls}">
        <div class="step-icon">{icon}</div>
        <div class="step-body">
            <div class="step-title">{title}</div>
            <div class="step-desc">{desc}</div>
        </div>
        <span class="step-badge {badge_cls}">{label}</span>
    </div>
    """, unsafe_allow_html=True)


def step_state(step_key, results, running):
    steps = ["search", "reader", "writer", "critic"]
    if step_key in results:
        return "done"
    if running:
        for k in steps:
            if k not in results:
                return "running" if k == step_key else "waiting"
    return "waiting"


def make_pdf(report_text: str, topic: str) -> bytes:
    """Convert the markdown report to a simple PDF using reportlab."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    import re

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter,
                            leftMargin=inch, rightMargin=inch,
                            topMargin=inch, bottomMargin=inch)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle('CustomTitle',
        parent=styles['Title'],
        fontSize=20, spaceAfter=12, textColor=colors.HexColor('#1a365d'))
    h2_style = ParagraphStyle('CustomH2',
        parent=styles['Heading2'],
        fontSize=13, spaceAfter=6, textColor=colors.HexColor('#2b6cb0'))
    body_style = ParagraphStyle('CustomBody',
        parent=styles['Normal'],
        fontSize=10, leading=16, spaceAfter=8)
    bold_style = ParagraphStyle('BoldBody',
        parent=body_style, fontName='Helvetica-Bold')

    story = []
    for line in report_text.split('\n'):
        line = line.strip()
        if not line:
            story.append(Spacer(1, 6))
        elif line.startswith('# '):
            story.append(Paragraph(line[2:], title_style))
        elif line.startswith('## '):
            story.append(Paragraph(line[3:], h2_style))
        elif line.startswith('**') and line.endswith('**'):
            story.append(Paragraph(line[2:-2], bold_style))
        else:
            # strip markdown links for PDF body
            clean = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', line)
            clean = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', clean)
            clean = re.sub(r'\*([^*]+)\*', r'<i>\1</i>', clean)
            story.append(Paragraph(clean, body_style))

    doc.build(story)
    return buf.getvalue()


def make_docx(report_text: str, topic: str) -> bytes:
    """Convert the markdown report to a DOCX using python-docx."""
    try:
        from docx import Document as DocxDocument
        from docx.shared import Pt, RGBColor
        import re

        doc = DocxDocument()
        style = doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)

        for line in report_text.split('\n'):
            line = line.strip()
            if not line:
                doc.add_paragraph()
            elif line.startswith('# '):
                p = doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', line[2:]),
                                  style='List Bullet')
            elif re.match(r'^\d+\. ', line):
                doc.add_paragraph(re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1',
                                         re.sub(r'^\d+\. ', '', line)),
                                  style='List Number')
            else:
                clean = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', line)
                clean = re.sub(r'\*\*([^*]+)\*\*', r'\1', clean)
                doc.add_paragraph(clean)

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    except ImportError:
        return None


# ── Session state init ─────────────────────────────────────────────────────────
for key, default in [("results", {}), ("running", False), ("done", False),
                     ("topic_val", ""), ("settings_open", False)]:
    if key not in st.session_state:
        st.session_state[key] = default


# ── Hero ───────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="hero">
    <div class="hero-badge">Multi-Agent AI System</div>
    <h1>Research<em>Orbit</em></h1>
    <p class="hero-sub">
        Four specialised agents collaborate — searching, scraping, writing,
        and critiquing — to produce a polished research report on any topic.
    </p>
</div>
<div class="rule"></div>
""", unsafe_allow_html=True)


# ── Layout ─────────────────────────────────────────────────────────────────────
col_left, col_gap, col_right = st.columns([5, 0.4, 4])

with col_left:
    # ── Topic input ──
    st.markdown('<div class="card card-blue">', unsafe_allow_html=True)
    topic = st.text_input(
        "Research Topic",
        placeholder="e.g. Quantum computing breakthroughs in 2025",
        key="topic_input",
    )

    # ── Settings row ──
    with st.expander("⚙  Research settings", expanded=False):
        col_a, col_b = st.columns(2)
        with col_a:
            depth = st.selectbox(
                "Research Depth",
                options=["Quick (3 sources)", "Standard (5 sources)", "Deep dive (8 sources)"],
                index=1,
                key="depth_select",
            )
        with col_b:
            length = st.selectbox(
                "Report Length",
                options=["Brief (300–500 words)", "Standard (700–1000 words)", "Detailed (1500–2500 words)"],
                index=1,
                key="length_select",
            )

    run_btn = st.button("⚡  Run Research Pipeline", use_container_width=True)
    st.markdown('</div>', unsafe_allow_html=True)

    # ── Example chips ──
    st.markdown("""
    <div style="display:flex;gap:0.5rem;flex-wrap:wrap;margin-top:0.3rem;margin-bottom:1.5rem;align-items:center;">
        <span style="font-family:'DM Mono',monospace;font-size:0.65rem;color:#3d4a56;letter-spacing:0.1em;">TRY →</span>
    """, unsafe_allow_html=True)
    for ex in ["LLM agents 2025", "CRISPR gene editing", "Fusion energy progress", "Climate tech startups"]:
        st.markdown(f"""
        <span style="background:rgba(255,255,255,0.03);border:1px solid rgba(255,255,255,0.06);
            border-radius:6px;padding:0.2rem 0.65rem;font-size:0.73rem;color:#536070;
            font-family:'Outfit',sans-serif;">{ex}</span>
        """, unsafe_allow_html=True)
    st.markdown("</div>", unsafe_allow_html=True)

with col_right:
    st.markdown('<div style="padding-top:0.3rem;">', unsafe_allow_html=True)
    st.markdown('<p style="font-family:\'DM Mono\',monospace;font-size:0.65rem;letter-spacing:0.2em;text-transform:uppercase;color:#536070;margin-bottom:0.8rem;">Pipeline</p>', unsafe_allow_html=True)

    r = st.session_state.results
    running = st.session_state.running

    render_step("🔍", "Search Agent",  "Gathers recent web information",         step_state("search", r, running))
    render_step("📄", "Reader Agent",  "Scrapes & extracts deep content",        step_state("reader", r, running))
    render_step("✍️", "Writer Chain",  "Drafts the full research report",        step_state("writer", r, running))
    render_step("🧐", "Critic Chain",  "Reviews accuracy, depth & structure",    step_state("critic", r, running))
    st.markdown('</div>', unsafe_allow_html=True)


# ── Parse settings helpers ────────────────────────────────────────────────────
def parse_depth(d: str) -> int:
    return {"Quick (3 sources)": 3, "Standard (5 sources)": 5, "Deep dive (8 sources)": 8}.get(d, 5)

def parse_length(l: str) -> str:
    return {"Brief (300–500 words)": "brief", "Standard (700–1000 words)": "standard",
            "Detailed (1500–2500 words)": "detailed"}.get(l, "standard")


# ── Trigger pipeline ──────────────────────────────────────────────────────────
if run_btn:
    if not topic.strip():
        st.warning("Please enter a research topic first.")
    else:
        st.session_state.results  = {}
        st.session_state.running  = True
        st.session_state.done     = False
        st.session_state.topic_val = topic
        st.rerun()

if st.session_state.running and not st.session_state.done:
    results   = {}
    topic_val = st.session_state.topic_val
    depth_val = parse_depth(st.session_state.get("depth_select", "Standard (5 sources)"))
    length_val= parse_length(st.session_state.get("length_select", "Standard (700–1000 words)"))

    with st.spinner("🔍  Search Agent is working…"):
        search_agent = build_search_agent(max_results=depth_val)
        sr = search_agent.invoke({
            "messages": [("user", f"Find recent, reliable and detailed information about: {topic_val}")]
        })
        results["search"] = sr["messages"][-1].content
        st.session_state.results = dict(results)

    with st.spinner("📄  Reader Agent is scraping top resources…"):
        reader_agent = build_reader_agent()
        rr = reader_agent.invoke({
            "messages": [("user",
                f"Based on the following search results about '{topic_val}', "
                f"pick the most relevant URL and scrape it for deeper content.\n\n"
                f"Search Results:\n{results['search'][:800]}"
            )]
        })
        results["reader"] = rr["messages"][-1].content
        st.session_state.results = dict(results)

    with st.spinner("✍️  Writer is drafting the report…"):
        research_combined = (
            f"SEARCH RESULTS:\n{results['search']}\n\n"
            f"DETAILED SCRAPED CONTENT:\n{results['reader']}"
        )
        writer_chain = build_writer_chain(report_length=length_val)
        results["writer"] = writer_chain.invoke({
            "topic": topic_val,
            "research": research_combined
        })
        st.session_state.results = dict(results)

    with st.spinner("🧐  Critic is reviewing the report…"):
        results["critic"] = critic_chain.invoke({"report": results["writer"]})
        st.session_state.results = dict(results)

    st.session_state.running = False
    st.session_state.done = True
    st.rerun()


# ── Results ────────────────────────────────────────────────────────────────────
r = st.session_state.results

if r:
    st.markdown('<div class="rule"></div>', unsafe_allow_html=True)
    st.markdown('<p style="font-family:\'DM Mono\',monospace;font-size:0.65rem;letter-spacing:0.2em;text-transform:uppercase;color:#536070;margin-bottom:1.2rem;">Results</p>', unsafe_allow_html=True)

    res_left, res_gap, res_right = st.columns([6, 0.3, 3.5])

    with res_left:
        # ── Raw outputs (collapsed) ──
        if "search" in r:
            with st.expander("🔍 Search results (raw)", expanded=False):
                st.markdown('<div class="result-box"><div class="section-label label-blue">Search Agent Output</div>', unsafe_allow_html=True)
                st.text(r["search"])
                st.markdown('</div>', unsafe_allow_html=True)

        if "reader" in r:
            with st.expander("📄 Scraped content (raw)", expanded=False):
                st.markdown('<div class="result-box"><div class="section-label label-blue">Reader Agent Output</div>', unsafe_allow_html=True)
                st.text(r["reader"])
                st.markdown('</div>', unsafe_allow_html=True)

        # ── Final report ──
        if "writer" in r:
            st.markdown('<div class="card"><div class="section-label label-purple">📝 Final Research Report</div>', unsafe_allow_html=True)
            st.markdown(r["writer"])
            st.markdown('</div>', unsafe_allow_html=True)

            # ── Export buttons ──
            st.markdown('<p style="font-family:\'DM Mono\',monospace;font-size:0.65rem;letter-spacing:0.2em;text-transform:uppercase;color:#536070;margin:1rem 0 0.5rem;">Export report</p>', unsafe_allow_html=True)
            exp_col1, exp_col2, exp_col3 = st.columns(3)

            topic_slug = st.session_state.topic_val.lower().replace(" ", "_")[:30]
            ts = int(time.time())

            with exp_col1:
                st.download_button(
                    label="⬇ Markdown (.md)",
                    data=r["writer"],
                    file_name=f"report_{topic_slug}_{ts}.md",
                    mime="text/markdown",
                )
            with exp_col2:
                pdf_bytes = make_pdf(r["writer"], st.session_state.topic_val)
                st.download_button(
                    label="⬇ PDF (.pdf)",
                    data=pdf_bytes,
                    file_name=f"report_{topic_slug}_{ts}.pdf",
                    mime="application/pdf",
                )
            with exp_col3:
                docx_bytes = make_docx(r["writer"], st.session_state.topic_val)
                if docx_bytes:
                    st.download_button(
                        label="⬇ Word (.docx)",
                        data=docx_bytes,
                        file_name=f"report_{topic_slug}_{ts}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                else:
                    st.caption("Install `python-docx` for DOCX export.")

    with res_right:
        if "critic" in r:
            st.markdown('<div class="card card-green"><div class="section-label label-green">🧐 Critic Feedback</div>', unsafe_allow_html=True)
            st.markdown(r["critic"])
            st.markdown('</div>', unsafe_allow_html=True)


# ── Footer ─────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="footer">
    ResearchOrbit · LangChain multi-agent pipeline · Streamlit
</div>
""", unsafe_allow_html=True)
